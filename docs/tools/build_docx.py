"""Build docs/RPMS_UI_Design_Documentation.docx from docs/RPMS_UI_Design_Documentation.md.

The Markdown file is the single source of the chapter text, so the .md and .docx never drift
apart. Supported Markdown: headings (# to ####), paragraphs, **bold**, *italic*, `code`,
bullet and numbered lists, pipe tables (cells may hold one image each), images with an
optional {width=N} in cm, caption lines ("*Figure 2.x: ...*" below an image or image table,
"Table 2.x: ..." above a table), [TOC]/[LOF]/[LOT] placeholders and <!-- pagebreak -->.

Formatting: A4, Times New Roman 12 pt, bold headings, justified body text, captions in the
form "Figure 2.x: caption", images kept on the same page as their caption, contents lists
and page numbers as Word fields (refreshed by docs/tools/update_docx_fields.ps1).

The script also checks that figure and table numbers run in sequence and that every
"Figure 2.x" / "Table 2.x" reference in the text points to a caption that exists.

Usage: python docs/tools/build_docx.py
"""
import re
import sys
from pathlib import Path

import pymupdf
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
DOCS = ROOT / "docs"
SRC = DOCS / "RPMS_UI_Design_Documentation.md"
OUT = DOCS / "RPMS_UI_Design_Documentation.docx"

CHAPTER = "2"
FONT = "Times New Roman"
TEXT_WIDTH_CM = 15.92          # A4 (21 cm) minus 2 x 2.54 cm margins
MAX_IMAGE_HEIGHT_CM = 20.5     # leaves room for the caption on the same page
PX_PER_CM = 1200 / TEXT_WIDTH_CM  # a 1200 px crop fills the text width

FIG_CAPTION = re.compile(rf"^\*(Figure {CHAPTER}\.(\d+): .+)\*$")
TAB_CAPTION = re.compile(rf"^(Table {CHAPTER}\.(\d+): .+)$")
IMAGE = re.compile(r"^!\[([^\]]*)\]\(([^)]+)\)(\{width=([\d.]+)\})?$")
REF = re.compile(rf"\b(Figures?|Tables?)\s+({CHAPTER}\.\d+(?:(?:,\s*|\s+and\s+|\s+to\s+|–)\s*{CHAPTER}\.\d+)*)")
INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|\*[^*\s][^*]*?\*)")


# --------------------------------------------------------------------------- styles

def set_font(style_or_run, size=None, bold=None, italic=None):
    font = style_or_run.font
    font.name = FONT
    if size:
        font.size = Pt(size)
    if bold is not None:
        font.bold = bold
    if italic is not None:
        font.italic = italic
    font.color.rgb = RGBColor(0, 0, 0)
    rpr = style_or_run.element.get_or_add_rPr() if hasattr(style_or_run.element, "get_or_add_rPr") \
        else style_or_run.element.rPr
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        fonts.set(qn(attr), FONT)


def setup_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    set_font(normal, 12)
    pf = normal.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.line_spacing = 1.5
    pf.space_after = Pt(6)

    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12), ("Heading 4", 12)):
        st = styles[name]
        set_font(st, size, bold=True, italic=(name == "Heading 4"))
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        st.paragraph_format.space_before = Pt(18 if name in ("Heading 1", "Heading 2") else 12)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.keep_with_next = True
        st.paragraph_format.line_spacing = 1.15

    for name in ("Figure Caption", "Table Caption"):
        st = styles.add_style(name, 1)  # paragraph style
        st.base_style = normal
        set_font(st, 11)
        st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        st.paragraph_format.line_spacing = 1.15
        st.paragraph_format.space_before = Pt(4)
        st.paragraph_format.space_after = Pt(12 if name == "Figure Caption" else 4)
        if name == "Table Caption":
            st.paragraph_format.keep_with_next = True

    st = styles.add_style("Table Text", 1)
    st.base_style = normal
    set_font(st, 10)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.line_spacing = 1.0
    st.paragraph_format.space_after = Pt(0)

    st = styles.add_style("List Item", 1)
    st.base_style = normal
    st.paragraph_format.left_indent = Cm(0.9)
    st.paragraph_format.first_line_indent = Cm(-0.6)
    st.paragraph_format.space_after = Pt(4)

    st = styles.add_style("Front Heading", 1)
    st.base_style = normal
    set_font(st, 14, bold=True)
    st.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    st.paragraph_format.space_before = Pt(12)
    st.paragraph_format.keep_with_next = True


def setup_page(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.54))
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_field(footer, "PAGE", "1")


# --------------------------------------------------------------------------- fields

def add_field(paragraph, instruction, placeholder=""):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    run._r.append(begin)
    run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    run._r.append(instr)
    run = paragraph.add_run()
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.append(sep)
    paragraph.add_run(placeholder)
    run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(end)


# --------------------------------------------------------------------------- inline text

def add_inline(paragraph, text, size=None):
    text = text.replace("&nbsp;", " ")
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Courier New"
            run.font.size = Pt((size or 12) - 1)
            continue
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(part)
        if size:
            run.font.size = Pt(size)


def add_caption(doc, text, style):
    p = doc.add_paragraph(style=style)
    label, rest = text.split(":", 1)
    run = p.add_run(label + ":")
    run.bold = True
    add_inline(p, rest)
    return p


# --------------------------------------------------------------------------- images

def image_size_cm(path, width_cm=None, max_width_cm=TEXT_WIDTH_CM, max_height_cm=MAX_IMAGE_HEIGHT_CM):
    pix = pymupdf.Pixmap(str(path))
    w_px, h_px = pix.width, pix.height
    if width_cm is None:
        width_cm = min(max_width_cm, w_px / PX_PER_CM)
    width_cm = min(width_cm, max_width_cm)
    height_cm = width_cm * h_px / w_px
    if height_cm > max_height_cm:
        height_cm = max_height_cm
        width_cm = height_cm * w_px / h_px
    return width_cm, height_cm


def add_image(doc, path, width_cm=None):
    w, h = image_size_cm(path, width_cm)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    p.add_run().add_picture(str(path), width=Cm(w), height=Cm(h))


# --------------------------------------------------------------------------- tables

def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:cantSplit")
    el.set(qn("w:val"), "true")
    tr_pr.append(el)


def fixed_layout(table):
    tbl_pr = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    tbl_pr.append(layout)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def column_widths(rows, total=TEXT_WIDTH_CM):
    ncols = len(rows[0])
    weights = []
    for c in range(ncols):
        lengths = [len(r[c]) for r in rows if c < len(r)]
        avg = sum(lengths) / len(lengths)
        weights.append(max(7.0, min(55.0, avg)))
    s = sum(weights)
    return [total * w / s for w in weights]


def add_text_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    size = 8.5 if ncols >= 6 else 10
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)
    widths = column_widths(rows)
    for r_idx, row_cells in enumerate(rows):
        row = table.rows[r_idx]
        if r_idx == 0:
            repeat_header(row)
        cant_split(row)
        for c_idx in range(ncols):
            cell = row.cells[c_idx]
            cell.width = Cm(widths[c_idx])
            text = row_cells[c_idx] if c_idx < len(row_cells) else ""
            p = cell.paragraphs[0]
            p.style = doc.styles["Table Text"]
            for k, piece in enumerate(text.split("<br>")):
                if k:
                    p = cell.add_paragraph(style="Table Text")
                add_inline(p, piece, size=size)
            if r_idx == 0:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
                shade(cell, "D9D9D9")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_image_table(doc, rows):
    header, images = rows[0], rows[1]
    ncols = len(header)
    col_w = TEXT_WIDTH_CM / ncols
    table = doc.add_table(rows=2, cols=ncols)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)
    for r in table.rows:
        cant_split(r)
    max_h = MAX_IMAGE_HEIGHT_CM - 1.5
    for c in range(ncols):
        hcell = table.rows[0].cells[c]
        hcell.width = Cm(col_w)
        p = hcell.paragraphs[0]
        p.style = doc.styles["Table Text"]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        add_inline(p, header[c], size=10)
        for run in p.runs:
            run.bold = True
        shade(hcell, "D9D9D9")

        icell = table.rows[1].cells[c]
        icell.width = Cm(col_w)
        m = IMAGE.match(images[c])
        if not m:
            raise ValueError(f"Expected an image in table cell: {images[c]}")
        path = DOCS / m.group(2)
        w, h = image_size_cm(path, col_w - 0.35, max_width_cm=col_w - 0.35, max_height_cm=max_h)
        p = icell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.line_spacing = 1.0
        p.add_run().add_picture(str(path), width=Cm(w), height=Cm(h))
    return table


# --------------------------------------------------------------------------- parser

def parse_and_build(doc, lines):
    figures, tables, refs = [], [], []
    i = 0
    paragraph_buf = []

    def flush_paragraph():
        if paragraph_buf:
            text = " ".join(s.strip() for s in paragraph_buf)
            refs.extend(find_refs(text))
            p = doc.add_paragraph()
            add_inline(p, text)
            paragraph_buf.clear()

    while i < len(lines):
        line = lines[i].rstrip("\n")
        stripped = line.strip()

        if not stripped:
            flush_paragraph()
            i += 1
            continue

        if stripped == "<!-- pagebreak -->":
            flush_paragraph()
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            i += 1
            continue

        if stripped in ("[TOC]", "[LOF]", "[LOT]"):
            flush_paragraph()
            title, instr = {
                "[TOC]": ("Contents", 'TOC \\o "1-3" \\h \\z \\u'),
                "[LOF]": ("List of Figures", 'TOC \\h \\z \\t "Figure Caption,1"'),
                "[LOT]": ("List of Tables", 'TOC \\h \\z \\t "Table Caption,1"'),
            }[stripped]
            doc.add_paragraph(title, style="Front Heading")
            p = doc.add_paragraph()
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_field(p, instr, "Right-click and choose Update Field to build this list.")
            i += 1
            continue

        heading = re.match(r"^(#{1,4})\s+(.*)$", stripped)
        if heading:
            flush_paragraph()
            level = len(heading.group(1))
            text = heading.group(2)
            if level == 1:
                doc.add_heading(text, level=1)
            else:
                doc.add_heading(text, level=level)
            i += 1
            continue

        m = FIG_CAPTION.match(stripped)
        if m:
            flush_paragraph()
            figures.append(int(m.group(2)))
            add_caption(doc, m.group(1), "Figure Caption")
            i += 1
            continue

        m = TAB_CAPTION.match(stripped)
        if m:
            flush_paragraph()
            tables.append(int(m.group(2)))
            add_caption(doc, m.group(1), "Table Caption")
            i += 1
            continue

        m = IMAGE.match(stripped)
        if m:
            flush_paragraph()
            path = DOCS / m.group(2)
            if not path.exists():
                raise FileNotFoundError(path)
            add_image(doc, path, float(m.group(4)) if m.group(4) else None)
            i += 1
            continue

        if stripped.startswith("|"):
            flush_paragraph()
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [split_row(r) for k, r in enumerate(block) if k != 1]
            for r in rows:
                refs.extend(find_refs(" ".join(r)))
            if any(IMAGE.match(c) for c in rows[-1]):
                add_image_table(doc, rows)
            else:
                add_text_table(doc, rows)
            continue

        bullet = re.match(r"^- (.*)$", stripped)
        number = re.match(r"^(\d+)\. (.*)$", stripped)
        if bullet or number:
            flush_paragraph()
            marker = "•" if bullet else number.group(1) + "."
            text = bullet.group(1) if bullet else number.group(2)
            # continuation lines that are indented belong to the item
            i += 1
            while i < len(lines) and lines[i].startswith("  ") and lines[i].strip() \
                    and not re.match(r"^\s*(- |\d+\. )", lines[i]):
                text += " " + lines[i].strip()
                i += 1
            refs.extend(find_refs(text))
            p = doc.add_paragraph(style="List Item")
            p.add_run(marker + "\t")
            p.paragraph_format.tab_stops.add_tab_stop(Cm(0.9))
            add_inline(p, text)
            continue

        paragraph_buf.append(stripped)
        i += 1

    flush_paragraph()
    return figures, tables, refs


def find_refs(text):
    found = []
    for m in REF.finditer(text):
        kind = "Figure" if m.group(1).startswith("Figure") else "Table"
        nums = [int(n.split(".")[1]) for n in re.findall(rf"{CHAPTER}\.\d+", m.group(2))]
        if " to " in m.group(2) and len(nums) == 2:
            nums = list(range(nums[0], nums[1] + 1))
        found.extend((kind, n) for n in nums)
    return found


def check_numbering(figures, tables, refs):
    problems = []
    for kind, nums in (("Figure", figures), ("Table", tables)):
        expected = list(range(1, len(nums) + 1))
        if nums != expected:
            problems.append(f"{kind} numbers are not in sequence: {nums}")
    fig_set, tab_set = set(figures), set(tables)
    for kind, n in refs:
        if n not in (fig_set if kind == "Figure" else tab_set):
            problems.append(f"Reference to missing {kind} {CHAPTER}.{n}")
    referenced = {(k, n) for k, n in refs}
    unref_f = [n for n in figures if ("Figure", n) not in referenced]
    unref_t = [n for n in tables if ("Table", n) not in referenced]
    return problems, unref_f, unref_t


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    setup_styles(doc)
    setup_page(doc)
    figures, tables, refs = parse_and_build(doc, lines)
    problems, unref_f, unref_t = check_numbering(figures, tables, refs)
    doc.save(OUT)
    print(f"Saved {OUT.relative_to(ROOT)}")
    print(f"Figures: {len(figures)}  Tables: {len(tables)}  Cross-references checked: {len(refs)}")
    if unref_f:
        print("Figures not referenced in the text (caption only):", ", ".join(f"{CHAPTER}.{n}" for n in unref_f))
    if unref_t:
        print("Tables not referenced in the text:", ", ".join(f"{CHAPTER}.{n}" for n in unref_t))
    if problems:
        print("PROBLEMS:")
        for p in problems:
            print("  -", p)
        sys.exit(1)
    print("Numbering and cross-references OK")


if __name__ == "__main__":
    main()
