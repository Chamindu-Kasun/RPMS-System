"""Structural check of the saved .docx (after Word has refreshed its fields).

- The three contents lists were filled in by Word (no placeholder text left).
- Counts the pictures in the document.
- Every figure caption directly follows a picture paragraph or a table of pictures, and
  every table caption directly precedes a table.

Usage: python docs/tools/check_docx.py
"""
import re
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
DOCX = ROOT / "docs" / "RPMS_UI_Design_Documentation.docx"


def main():
    doc = Document(DOCX)
    body = doc.element.body
    blocks = [el for el in body.iterchildren() if el.tag in (qn("w:p"), qn("w:tbl"))]
    pictures = len(body.xpath(".//*[local-name()='pic']"))
    text_of = lambda el: "".join(t.text or "" for t in el.iter(qn("w:t")))
    has_pic = lambda el: bool(el.xpath(".//*[local-name()='pic']"))

    problems = []
    toc_entries = {"Figure": 0, "Table": 0, "Heading": 0}
    placeholders = 0
    figures = tables = 0
    for k, el in enumerate(blocks):
        if el.tag != qn("w:p"):
            continue
        text = text_of(el)
        style = el.find(qn("w:pPr") + "/" + qn("w:pStyle"))
        style = style.get(qn("w:val")) if style is not None else ""
        if "Update Field" in text:
            placeholders += 1
        if style.startswith("TOC") or style.startswith("TableofFigures"):
            if text.startswith("Figure 2."):
                toc_entries["Figure"] += 1
            elif text.startswith("Table 2."):
                toc_entries["Table"] += 1
            else:
                toc_entries["Heading"] += 1
            continue
        if re.match(r"^Figure 2\.\d+:", text) and style == "FigureCaption":
            figures += 1
            prev = blocks[k - 1]
            if not has_pic(prev):
                problems.append(f"'{text[:40]}' does not follow a picture")
        if re.match(r"^Table 2\.\d+:", text) and style == "TableCaption":
            tables += 1
            if blocks[k + 1].tag != qn("w:tbl"):
                problems.append(f"'{text[:40]}' is not followed by a table")

    print(f"Pictures: {pictures}")
    print(f"Figure captions: {figures}  Table captions: {tables}")
    print(f"Contents entries: {toc_entries}")
    if placeholders:
        problems.append(f"{placeholders} field placeholder(s) left")
    print("PROBLEMS:\n  - " + "\n  - ".join(problems) if problems else "DOCX check OK")


if __name__ == "__main__":
    main()
